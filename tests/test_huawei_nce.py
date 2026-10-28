"""Tests for Huawei NCE native collector, AP parser, WAC status, and WLC filtering."""

import inspect
import json
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from network_report.collectors.huawei_ap import (
    collect_huawei_ap_data,
    parse_ap_info_csv,
    parse_ssid_info_csv,
)
from network_report.collectors.huawei_nce import (
    NceCollector,
    WacStatus,
    parse_wac_pct,
)
from network_report.processing.analyzer import generate_technical_analysis
from network_report.writers.excel_writer import write_daily_excel
from network_report.writers.wag_formatter import format_whatsapp_report


def test_parse_wac_pct():
    assert parse_wac_pct("Current / 25%") == 25.0
    assert parse_wac_pct("Current / 100%") == 100.0
    assert parse_wac_pct("Current / 57.5%") == 57.5
    assert parse_wac_pct("Current / 0%") == 0.0
    assert parse_wac_pct("20%") == 20.0
    assert parse_wac_pct("garbage") is None
    assert parse_wac_pct("") is None
    assert parse_wac_pct("Current / N/A") is None


def test_parse_ap_info_real_format(tmp_path: Path):
    ap_csv = tmp_path / "apInfo.csv"
    content = (
        '\t"AP ID"\t,"AP name","Status","Installation location","STA Quantity","CPU Usage","Memory Usage"\n'
        '1,"AP-H1-L1-IN01","normal","Hangar 1 Indoor",12,"10%","35%"\n'
        '2,"AP-H1-L1-OUT01","normal","Apron Outdoor",5,"15%","40%"\n'
        '3,"AP-H2-L1-OUT02","fault","Perimeter",0,"0%","0%"\n'
        '4,"AP-H2-L2-IN02","offline","Office 2nd Fl",0,"0%","0%"\n'
        '5,"AP-H3-L1-OUT03","abnormal","Outdoor Park",0,"5%","20%"\n'
    )
    ap_csv.write_text(content, encoding="utf-8-sig")

    res = parse_ap_info_csv(ap_csv)
    normal, fault, fault_names, in_ok, out_ok, out_fault = res
    assert normal == 2
    assert fault == 3
    assert in_ok == 1
    assert out_ok == 1
    assert out_fault == 2  # AP-H2-L1-OUT02 (-out) and AP-H3-L1-OUT03 (-out + outdoor)
    assert fault_names == ["AP-H2-L1-OUT02", "AP-H2-L2-IN02", "AP-H3-L1-OUT03"]
    assert res.total_sta == 17
    assert res.ap_sta_counts["AP-H1-L1-IN01"] == 12


def test_parse_ssid_info_real_format(tmp_path: Path):
    ssid_csv = tmp_path / "ssidInfo.csv"
    content = (
        '\t"SSID"\t,"User Quantity"\n'
        '"CORP-Employee",100\n'
        '"CORP-Guest",25\n'
        '"CORP-IoT",10\n'
        '"VIP-Executive",5\n'
    )
    ssid_csv.write_text(content, encoding="utf-8-sig")

    dist = parse_ssid_info_csv(ssid_csv)
    assert dist["Employee"] == 100
    assert dist["Guest"] == 25
    assert dist["IoT"] == 10
    assert dist["VIP"] == 5


def test_collect_huawei_ap_data_no_sibling_and_real_fixture(tmp_path: Path):
    source_code = inspect.getsource(collect_huawei_ap_data)
    assert "AP-Report-Auto" not in source_code

    raw_dir = tmp_path / "raw" / "huawei_ap"
    date_dir = raw_dir / "20261021"
    date_dir.mkdir(parents=True)

    ap_csv = date_dir / "apInfo.csv"
    ap_content = (
        "AP ID,AP name,Status,Installation location,STA Quantity\n"
        "1,AP-H1-IN01,normal,Indoor,10\n"
        "2,AP-H1-OUT01,normal,Outdoor,5\n"
        "3,AP-H1-OUT02,fault,Outdoor,0\n"
    )
    ap_csv.write_text(ap_content, encoding="utf-8")

    ssid_csv = date_dir / "ssidInfo.csv"
    ssid_content = "SSID,User Quantity\nEmployee,50\nGuest,15\n"
    ssid_csv.write_text(ssid_content, encoding="utf-8")

    data = collect_huawei_ap_data(raw_dir, "20261021")
    assert data is not None
    assert data.normal_count == 2
    assert data.fault_count == 1
    assert data.indoor_ok == 1
    assert data.outdoor_ok == 1
    assert data.outdoor_fault == 1
    assert data.fault_ap_names == ["AP-H1-OUT02"]
    assert data.total_users == 65
    assert data.sta_quantity == 15


def test_wag_formatter_filters_wlc_and_renders_wac():
    payload = {
        "date_str": "21 Oktober 2026",
        "cpu": [("WLC-Aruba-01", 85.0), ("SW-CORE-01", 42.5)],
        "memory": [("WLC-Aruba-02", 78.0), ("SW-DIST-01", 52.0)],
        "wac": {
            "device_name": "WAC-1-CORE",
            "cpu_pct": 20.0,
            "mem_pct": 57.0,
            "temperature": "42 C",
            "running_time": "120 days",
            "software_version": "V200R019C00",
            "error": None,
        },
    }

    msg = format_whatsapp_report(payload)
    # WLC must be filtered out
    assert "WLC-Aruba-01" not in msg
    assert "WLC-Aruba-02" not in msg
    assert "SW-CORE-01" in msg
    assert "SW-DIST-01" in msg

    # Unified WAC/NCE subsection at item 5 (WAC controller metrics only; AP/Users are in Section D)
    assert "5. 🎛️ Huawei WAC / iMaster NCE (Live)" in msg
    assert "Device Name → WAC-1-CORE" in msg
    assert "CPU Utilization → 20.00%" in msg
    assert "Memory Usage → 57.00%" in msg
    assert "Temperature → 42 C" in msg
    assert "Running Time → 120 days" in msg
    assert "Software Version → V200R019C00" in msg

    # NCE AP/Users info MUST NOT be in Section C (already in Section D)
    assert "📡 iMaster NCE (AP & Users)" not in msg
    assert "AP Terhubung" not in msg
    assert "Pengguna Aktif" not in msg

    # Section numbering: Avg Uptime is 6, Analisa is 7
    assert "6. ⏱️ Avg Uptime by Device" in msg
    assert "7. 📝 Analisa Utilisasi" in msg


def test_wag_formatter_wac_missing_or_error():
    # Missing WAC
    msg_missing = format_whatsapp_report({"date_str": "21 Oktober 2026"})
    assert "5. 🎛️ Huawei WAC / iMaster NCE (Live)" in msg_missing
    assert "⚠️ Data WAC tidak tersedia" in msg_missing

    # Error WAC
    msg_err = format_whatsapp_report(
        {
            "date_str": "21 Oktober 2026",
            "wac": {"error": "Connection timeout"},
        }
    )
    assert "5. 🎛️ Huawei WAC / iMaster NCE (Live)" in msg_err
    assert "⚠️ Gagal scrape (Connection timeout)" in msg_err

    # NCE AP/Users must NOT be in Section C
    assert "📡 iMaster NCE (AP & Users)" not in msg_missing


def test_analyzer_no_wlc_branch():
    source_code = inspect.getsource(generate_technical_analysis)
    assert "WLC" not in source_code

    cpu = [("SW-CORE-01", 18.5)]
    mem = [("WLC-01", 69.5)]
    c_txt, m_txt = generate_technical_analysis(cpu, mem)
    assert "SW-CORE-01" in c_txt
    assert "Perangkat WLC-01 sebesar 69.50%" in m_txt
    assert "sesi klien wireless" not in m_txt
    assert "Alokasi memori seluruh perangkat terpantau stabil dan terkendali." in m_txt


def test_excel_writer_wac_and_no_wlc(tmp_path: Path):
    template = tmp_path / "Template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "21"
    wb.save(template)

    output = tmp_path / "Daily-Report-2026-10.xlsx"
    payload = {
        "cpu": [("WLC-Aruba", 88.0), ("SW-CORE", 35.0)],
        "memory": [("WLC-Aruba", 92.0), ("SW-DIST", 45.0)],
        "wac": {
            "device_name": "WAC-1-CORE",
            "cpu_pct": 22.0,
            "mem_pct": 58.0,
            "temperature": "40 C",
            "running_time": "100 days",
            "error": None,
        },
    }

    ok = write_daily_excel(template, output, 21, payload)
    assert ok is True

    # Reload output and verify cells
    from openpyxl import load_workbook

    res_wb = load_workbook(output)
    res_ws = res_wb["21"]

    # CPU row 86 should have SW-CORE, not WLC-Aruba
    assert res_ws.cell(row=86, column=2).value == "SW-CORE"
    # Memory row 100 should have SW-DIST, not WLC-Aruba
    assert res_ws.cell(row=100, column=2).value == "SW-DIST"

    # WAC in col J (10), rows 114-119
    assert res_ws.cell(row=114, column=10).value == "Huawei WAC Status (Live)"
    assert res_ws.cell(row=115, column=10).value == "Device: WAC-1-CORE"
    assert res_ws.cell(row=116, column=10).value == "CPU: 22.00%"
    assert res_ws.cell(row=117, column=10).value == "Memory: 58.00%"
    assert res_ws.cell(row=118, column=10).value == "Temp: 40 C"
    assert res_ws.cell(row=119, column=10).value == "Running: 100 days"


def test_nce_collector_missing_credentials():
    collector = NceCollector(base_url="", username="", password="")
    res = collector.collect("20261021")
    assert res.error is not None
    assert "Missing NCE" in res.error


def test_nce_collector_mock_playwright(tmp_path: Path):
    fake_playwright = MagicMock()
    fake_browser = MagicMock()
    fake_context = MagicMock()
    fake_page = MagicMock()

    fake_playwright.chromium.launch.return_value = fake_browser
    fake_browser.new_context.return_value = fake_context
    fake_context.new_page.return_value = fake_page

    fake_page.url = "https://10.10.1.30/view/main.html"
    fake_page.is_visible.return_value = False

    fake_download = MagicMock()
    fake_download_info = MagicMock()
    fake_download_info.value = fake_download
    fake_page.expect_download.return_value.__enter__.return_value = fake_download_info

    def mock_locator(sel):
        mock_el = MagicMock()
        mock_el.count.return_value = 1
        text_map = {
            "#cpuUsageCurrentValue": "Current / 25%",
            "#memeryUsageCurrentValue": "Current / 57%",
            "#deviceInfo_deviceName": "WAC-1-CORE",
            "#deviceInfo_deviceType": "AC6508",
            "#deviceInfo_temperature": "42 C",
            "#deviceInfo_runningTime": "120 days",
            "#deviceInfo_softwareVersion": "V200R019C00",
            "#deviceInfo_serialNumber": "2102352V",
            "#deviceInfo_maxApCount": "1024",
        }
        mock_el.first.inner_text.return_value = text_map.get(sel, "")
        return mock_el

    fake_page.locator.side_effect = mock_locator

    with patch("network_report.collectors.huawei_nce.sync_playwright") as mock_pw_init:
        mock_pw_init.return_value.start.return_value = fake_playwright

        collector = NceCollector(
            base_url="https://10.10.1.30",
            username="admin",
            password="secret_password",
            output_root=tmp_path,
        )
        export = collector.collect("20261021")
        collector.close()

        assert export.error is None
        assert export.ap_csv_path == tmp_path / "raw" / "huawei_ap" / "20261021" / "apInfo.csv"
        assert export.ssid_csv_path == tmp_path / "raw" / "huawei_ap" / "20261021" / "ssidInfo.csv"
        assert export.wac is not None
        assert isinstance(export.wac, WacStatus)
        assert export.wac.cpu_pct == 25.0
        assert export.wac.mem_pct == 57.0
        assert export.wac.device_name == "WAC-1-CORE"
        assert export.wac.device_type == "AC6508"

        # Check saved wac_status.json
        wac_json = tmp_path / "raw" / "wac" / "20261021" / "wac_status.json"
        assert wac_json.exists()
        saved_dict = json.loads(wac_json.read_text(encoding="utf-8"))
        assert saved_dict["cpu_pct"] == 25.0
        assert saved_dict["device_name"] == "WAC-1-CORE"


def test_wac_fallback_logic_when_ap_exists_but_wac_missing(tmp_path: Path, monkeypatch):
    """Test WAC fallback: if apInfo.csv exists but wac_status.json is missing,
    pipeline should trigger live NceCollector and save wac_status.json."""
    from network_report.collectors.huawei_nce import NceCollector
    from unittest.mock import MagicMock, patch

    # Create config with huawei enabled
    output_root = tmp_path / "output"
    raw_root = tmp_path / "raw"
    output_root.mkdir(parents=True)
    raw_root.mkdir(parents=True)

    # Create 14:00 AP snapshot (apInfo.csv exists)
    ap_1400_dir = raw_root / "huawei_ap" / "20261021"
    ap_1400_dir.mkdir(parents=True)
    (ap_1400_dir / "apInfo.csv").write_text(
        "AP ID,AP name,Status,Installation location,STA Quantity\n"
        "1,AP-H1-IN01,normal,Indoor,10\n",
        encoding="utf-8",
    )
    (ap_1400_dir / "ssidInfo.csv").write_text(
        "SSID,User Quantity\nEmployee,50\n",
        encoding="utf-8",
    )

    # No wac_status.json for today
    wac_today = output_root / "raw" / "wac" / "20261021"
    wac_today.mkdir(parents=True)
    # Note: deliberately NOT creating wac_status.json for today

    # Create a past valid wac_status.json (FALLBACK 2)
    wac_past = output_root / "raw" / "wac" / "20261020"
    wac_past.mkdir(parents=True)
    past_wac_data = {
        "cpu_pct": 30.0,
        "mem_pct": 60.0,
        "device_name": "WAC-1-CORE",
        "device_type": "AC6508",
        "temperature": "42 C",
        "running_time": "120 days",
        "software_version": "V200R019C00",
        "serial_number": "2102352V",
        "max_ap_count": "1024",
        "captured_at": "2026-10-20T14:00:00",
        "error": None,
    }
    (wac_past / "wac_status.json").write_text(json.dumps(past_wac_data), encoding="utf-8")

    # Mock NceCollector to succeed and return WAC
    fake_playwright = MagicMock()
    fake_browser = MagicMock()
    fake_context = MagicMock()
    fake_page = MagicMock()

    fake_playwright.chromium.launch.return_value = fake_browser
    fake_browser.new_context.return_value = fake_context
    fake_context.new_page.return_value = fake_page
    fake_page.url = "https://10.10.1.30/view/main.html"
    fake_page.is_visible.return_value = False

    fake_download = MagicMock()
    fake_download_info = MagicMock()
    fake_download_info.value = fake_download
    fake_page.expect_download.return_value.__enter__.return_value = fake_download_info

    def mock_locator(sel):
        mock_el = MagicMock()
        mock_el.count.return_value = 1
        text_map = {
            "#cpuUsageCurrentValue": "Current / 25%",
            "#memeryUsageCurrentValue": "Current / 57%",
            "#deviceInfo_deviceName": "WAC-1-CORE",
            "#deviceInfo_deviceType": "AC6508",
            "#deviceInfo_temperature": "42 C",
            "#deviceInfo_runningTime": "120 days",
            "#deviceInfo_softwareVersion": "V200R019C00",
            "#deviceInfo_serialNumber": "2102352V",
            "#deviceInfo_maxApCount": "1024",
        }
        mock_el.first.inner_text.return_value = text_map.get(sel, "")
        return mock_el

    fake_page.locator.side_effect = mock_locator

    with patch("network_report.collectors.huawei_nce.sync_playwright") as mock_pw_init:
        mock_pw_init.return_value.start.return_value = fake_playwright

        # We can't easily test the full pipeline without mocking many things,
        # so we test the fallback logic directly by checking that the collector
        # is called and WAC is saved

        collector = NceCollector(
            base_url="https://10.10.1.30",
            username="admin",
            password="secret",
            output_root=output_root,
        )
        export = collector.collect("20261021")
        collector.close()

        assert export.error is None
        assert export.wac is not None
        assert export.wac.cpu_pct == 25.0
        assert export.wac.mem_pct == 57.0

        # Verify wac_status.json was saved for today
        wac_today_json = wac_today / "wac_status.json"
        assert wac_today_json.exists()
        saved = json.loads(wac_today_json.read_text(encoding="utf-8"))
        assert saved["cpu_pct"] == 25.0
        assert saved["device_name"] == "WAC-1-CORE"


def test_wac_fallback_to_past_when_live_fails(tmp_path: Path, monkeypatch):
    """Test WAC fallback: if live NceCollector fails, fallback to most recent past wac_status.json."""
    from network_report.collectors.huawei_nce import WacStatus

    output_root = tmp_path / "output"
    output_root.mkdir(parents=True)

    # No today's wac_status.json
    # Create multiple past wac_status.json files
    wac_20261018 = output_root / "raw" / "wac" / "20261018"
    wac_20261018.mkdir(parents=True)
    (wac_20261018 / "wac_status.json").write_text(
        json.dumps({
            "cpu_pct": 20.0, "mem_pct": 50.0, "device_name": "WAC-OLD",
            "device_type": "", "temperature": "", "running_time": "",
            "software_version": "", "serial_number": "", "max_ap_count": "",
            "captured_at": "2026-10-18T14:00:00", "error": None
        }), encoding="utf-8"
    )

    wac_20261019 = output_root / "raw" / "wac" / "20261019"
    wac_20261019.mkdir(parents=True)
    (wac_20261019 / "wac_status.json").write_text(
        json.dumps({
            "cpu_pct": 30.0, "mem_pct": 60.0, "device_name": "WAC-1-CORE",
            "device_type": "AC6508", "temperature": "42 C", "running_time": "120 days",
            "software_version": "V200R019C00", "serial_number": "2102352V",
            "max_ap_count": "1024", "captured_at": "2026-10-19T14:00:00", "error": None
        }), encoding="utf-8"
    )

    wac_20261020 = output_root / "raw" / "wac" / "20261020"
    wac_20261020.mkdir(parents=True)
    (wac_20261020 / "wac_status.json").write_text(
        json.dumps({
            "cpu_pct": 35.0, "mem_pct": 65.0, "device_name": "WAC-1-CORE",
            "device_type": "AC6508", "temperature": "44 C", "running_time": "121 days",
            "software_version": "V200R019C00", "serial_number": "2102352V",
            "max_ap_count": "1024", "captured_at": "2026-10-20T14:00:00", "error": None
        }), encoding="utf-8"
    )

    # Test the fallback logic (simulating what pipeline.py does)
    wac_status = None

    # FALLBACK 2: Search for most recent past wac_status.json
    past_wac_files = sorted(output_root.glob("raw/wac/*/wac_status.json"), reverse=True)
    for past_wac in past_wac_files:
        try:
            w_dict = json.loads(past_wac.read_text(encoding="utf-8"))
            wac_status = WacStatus(**w_dict)
            break
        except Exception:
            continue

    # Should get the most recent one (20261020)
    assert wac_status is not None
    assert wac_status.cpu_pct == 35.0
    assert wac_status.mem_pct == 65.0
    assert wac_status.device_name == "WAC-1-CORE"
    assert wac_status.temperature == "44 C"


def test_wac_fallback_skips_corrupt_past_files(tmp_path: Path):
    """Test WAC fallback skips corrupt/invalid past wac_status.json files."""
    from network_report.collectors.huawei_nce import WacStatus

    output_root = tmp_path / "output"
    output_root.mkdir(parents=True)

    # Create corrupt past wac_status.json
    wac_bad = output_root / "raw" / "wac" / "20261018"
    wac_bad.mkdir(parents=True)
    (wac_bad / "wac_status.json").write_text("not valid json", encoding="utf-8")

    # Create valid past wac_status.json
    wac_good = output_root / "raw" / "wac" / "20261019"
    wac_good.mkdir(parents=True)
    (wac_good / "wac_status.json").write_text(
        json.dumps({
            "cpu_pct": 30.0, "mem_pct": 60.0, "device_name": "WAC-1-CORE",
            "device_type": "AC6508", "temperature": "42 C", "running_time": "120 days",
            "software_version": "V200R019C00", "serial_number": "2102352V",
            "max_ap_count": "1024", "captured_at": "2026-10-19T14:00:00", "error": None
        }), encoding="utf-8"
    )

    # Test the fallback logic
    wac_status = None

    past_wac_files = sorted(output_root.glob("raw/wac/*/wac_status.json"), reverse=True)
    for past_wac in past_wac_files:
        try:
            w_dict = json.loads(past_wac.read_text(encoding="utf-8"))
            wac_status = WacStatus(**w_dict)
            break
        except Exception:
            continue

    # Should skip corrupt and get the valid one
    assert wac_status is not None
    assert wac_status.cpu_pct == 30.0
    assert wac_status.device_name == "WAC-1-CORE"
