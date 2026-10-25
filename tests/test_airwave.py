"""Unit tests for native AirWave collector & client (Phase 8B)."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl

from network_report.collectors.airwave import (
    AirWaveClient,
    build_report_window,
    collect_airwave_data,
    extract_report_id,
    is_wlc_device,
    parse_top_csv,
)
from network_report.writers.excel_writer import write_daily_excel
from network_report.writers.wag_formatter import format_whatsapp_report


def test_build_report_window():
    d = date(2026, 10, 19)
    start_str, end_str = build_report_window(d)
    assert start_str == "10/19/2026 00:00"
    assert end_str == "10/19/2026 23:59"


def test_extract_report_id():
    # Completed row with ID
    html_ok = (
        "<table><tr><td>10/19/2026</td>"
        '<td><a href="/report_detail?id=5356">#UP time Report Switch Daily</a></td>'
        "<td>Completed</td></tr></table>"
    )
    assert extract_report_id(html_ok) == "5356"

    # Pending row -> None
    html_pending = (
        "<table><tr><td>10/19/2026</td>"
        '<td><a href="/report_detail?id=5356">#UP time Report Switch Daily</a></td>'
        "<td>Pending in queue</td></tr></table>"
    )
    assert extract_report_id(html_pending) is None

    # In Progress row -> None
    html_in_progress = (
        "<table><tr><td>10/19/2026</td>"
        '<td><a href="/report_detail?id=5356">#UP time Report Switch Daily</a></td>'
        "<td>In Progress</td></tr></table>"
    )
    assert extract_report_id(html_in_progress) is None

    # No report_detail link -> None
    html_no_link = "<table><tr><td>10/19/2026</td><td>No reports</td></tr></table>"
    assert extract_report_id(html_no_link) is None


def test_is_wlc_device():
    assert is_wlc_device("WLC-MD-L2-7210") is True
    assert is_wlc_device("wlc-backup-01") is True
    assert is_wlc_device("SW-CORE-01", "Wireless Controller") is True
    assert is_wlc_device("SW-CORE-01", "Switch") is False


def test_parse_top_csv_excludes_wlc(tmp_path: Path):
    csv_file = tmp_path / "top_memory_usage.csv"
    csv_file.write_text(
        "Device,Usage\n"
        "WLC-MD-L2-7210,69.64%\n"
        "ASW01-WS1-LGL2A-2530,43.85%\n"
        "wlc-secondary,55.00%\n"
        "ESW01-WS1-LGL2B-2530,43.78%\n"
    )

    results = parse_top_csv(csv_file)
    dev_names = [dev for dev, _ in results]
    assert "WLC-MD-L2-7210" not in dev_names
    assert "wlc-secondary" not in dev_names
    assert dev_names == ["ASW01-WS1-LGL2A-2530", "ESW01-WS1-LGL2B-2530"]
    assert results[0] == ("ASW01-WS1-LGL2A-2530", 43.85)


def test_collect_airwave_data_no_silent_fallback(tmp_path: Path):
    day18 = tmp_path / "20261018"
    day18.mkdir(parents=True)
    (day18 / "top_cpu_utilization.csv").write_text("SW-01,10.0%\n")

    # Day 19 does not exist -> must return None, not day 18
    data = collect_airwave_data(tmp_path, "20261019")
    assert data is None


def test_collect_airwave_data_wlc_excluded_from_uptime(tmp_path: Path):
    day_dir = tmp_path / "20261019"
    day_dir.mkdir(parents=True)

    (day_dir / "top_cpu_utilization_123.csv").write_text(
        "Device,Utilization\nSW-01,15.0%\nWLC-01,80.0%\n"
    )
    (day_dir / "top_memory_usage_123.csv").write_text("Device,Usage\nSW-01,40.0%\nWLC-01,75.0%\n")
    (day_dir / "avg_uptime_by_device_123.csv").write_text(
        "Device,Group,Device Uptime,SNMP Comm Uptime,ICMP Comm Uptime,HTTPS Comm Uptime\n"
        "SW-01,Switch,100.0%,100.0%,100.0%,-\n"
        "SW-02,Switch,98.0%,98.0%,98.0%,-\n"
        "WLC-MD-L2-7210,Wireless Controller,0.0%,0.0%,0.0%,-\n"
    )

    data = collect_airwave_data(tmp_path, "20261019")
    assert data is not None
    assert len(data.top_cpu) == 1
    assert data.top_cpu[0] == ("SW-01", 15.0)
    assert len(data.top_memory) == 1
    assert data.top_memory[0] == ("SW-01", 40.0)
    # Average of SW-01 (100) and SW-02 (98) is 99.0%, WLC (0%) must be excluded
    assert data.device_uptime_pct == 99.0
    assert data.snmp_uptime_pct == 99.0
    assert data.icmp_uptime_pct == 99.0


def test_collect_airwave_data_fallback_to_any_switch(tmp_path: Path):
    day_dir = tmp_path / "20261019"
    day_dir.mkdir(parents=True)

    (day_dir / "top_cpu_utilization_123.csv").write_text("Device,Utilization\nSW-01,15.0%\n")
    (day_dir / "top_memory_usage_123.csv").write_text("Device,Usage\nSW-01,40.0%\n")
    # avg_uptime_by_device is missing, avg_uptime_by_any_switch is present
    (day_dir / "avg_uptime_by_any_switch_123.csv").write_text(
        '"Device Uptime","SNMP Comm Uptime","ICMP Comm Uptime","HTTPS Comm Uptime"\n'
        '98.66%,98.66%,99.33%,"	-"\n'
    )

    data = collect_airwave_data(day_dir, "20261019")
    assert data is not None
    assert data.device_uptime_pct == 98.66
    assert data.snmp_uptime_pct == 98.66
    assert data.icmp_uptime_pct == 99.33


def test_wag_formatter_renders_airwave_failure_and_source_date():
    payload = {
        "date_str": "19 Oktober 2026",
        "airwave_error": "Connection refused",
        "airwave_source_date": "19 Oktober 2026",
        "iot": {"status": "Normal"},
        "prtg": {"up": 100, "paused": 0, "total": 100},
        "ping": [],
        "ap": {},
        "mrtg": {},
        "cpu": [],
        "memory": [],
        "uptime": {
            "device": None,
            "snmp": None,
            "icmp": None,
            "https": "-",
            "error": "Connection refused",
            "source_date": "19 Oktober 2026",
        },
        "orbit": [],
        "highlights": {"high_traffic": [], "hardware_fault": "None", "problem_open": "-"},
    }

    wag = format_whatsapp_report(payload)
    assert "Data: 19 Oktober 2026" in wag
    assert "⚠️ Gagal scrape" in wag


def test_excel_writer_renders_airwave_failure_and_source_date(tmp_path: Path):
    template = Path("config/Template-Daily-Report.xlsx")
    if not template.exists():
        template = Path("config/Template-Daily-Report.example.xlsx")
    out_file = tmp_path / "Daily-Report-2026-10.xlsx"

    payload = {
        "date_str": "19 Oktober 2026",
        "airwave_error": "Timeout polling report",
        "airwave_source_date": "19 Oktober 2026",
        "iot": {"status": "Normal"},
        "prtg": {"up": 100, "paused": 0, "total": 100},
        "ping": [],
        "ap": {},
        "mrtg": {},
        "cpu": [],
        "memory": [],
        "uptime": {
            "device": None,
            "snmp": None,
            "icmp": None,
            "https": "-",
            "error": "Timeout polling report",
            "source_date": "19 Oktober 2026",
        },
        "orbit": [],
        "highlights": {"high_traffic": [], "hardware_fault": "None", "problem_open": "-"},
    }

    ok = write_daily_excel(template, out_file, day_number=19, payload=payload)
    assert ok is True

    wb = openpyxl.load_workbook(out_file, data_only=True)
    ws = wb["19"]

    # Header date labels
    assert "Data: 19 Oktober 2026" in str(ws.cell(84, 1).value)
    assert "Data: 19 Oktober 2026" in str(ws.cell(98, 1).value)
    assert "Data: 19 Oktober 2026" in str(ws.cell(112, 1).value)

    # Status cells show failure
    assert ws.cell(86, 6).value == "⚠️ Gagal scrape"
    assert ws.cell(100, 6).value == "⚠️ Gagal scrape"
    assert ws.cell(114, 7).value == "⚠️ Gagal scrape"


def test_airwave_client_login_and_reuse(tmp_path: Path):
    session_file = tmp_path / "storage_state.json"
    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=session_file,
    )

    mock_sess = MagicMock()
    mock_sess.__enter__.return_value = mock_sess
    mock_sess.__exit__.return_value = None
    mock_post_resp = MagicMock()
    mock_post_resp.raise_for_status.return_value = None
    mock_cookie = MagicMock()
    mock_cookie.name = "MercuryAuthHandlerCookie_AMPAuth"
    mock_cookie.value = "auth-token-12345"
    mock_sess.cookies = [mock_cookie]

    with patch("network_report.collectors.airwave.requests.Session", return_value=mock_sess):
        client.login()

    assert client._auth_cookie == "auth-token-12345"
    assert session_file.exists()
    state = json.loads(session_file.read_text())
    assert state["cookies"][0]["value"] == "auth-token-12345"

    # Test reuse: when session_file exists and /nf/reports_list returns 200 with table marker
    mock_test_sess = MagicMock()
    mock_test_sess.__enter__.return_value = mock_test_sess
    mock_test_sess.__exit__.return_value = None
    mock_root_resp = MagicMock()
    mock_root_resp.status_code = 200
    mock_root_resp.text = (
        '<table><tr><td><a href="/report_detail?id=1">#UP time Report</a></td></tr></table>'
    )
    mock_test_sess.get.return_value = mock_root_resp

    client2 = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=session_file,
    )

    with patch("network_report.collectors.airwave.requests.Session", return_value=mock_test_sess):
        client2.login()

    assert client2._auth_cookie == "auth-token-12345"
    mock_test_sess.get.assert_called_with("https://10.10.1.20/nf/reports_list", timeout=15)
    mock_test_sess.post.assert_not_called()


def test_airwave_client_fetch_report_mocked(tmp_path: Path):
    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=tmp_path / "session.json",
        poll_interval_seconds=0,
        poll_timeout_minutes=1,
    )
    client._auth_cookie = "mock-cookie-val"

    # Mock playwright
    mock_playwright = MagicMock()
    mock_browser = MagicMock()
    mock_context = MagicMock()
    mock_page = MagicMock()

    mock_playwright.firefox.launch.return_value = mock_browser
    mock_browser.new_context.return_value = mock_context
    mock_context.new_page.return_value = mock_page

    mock_locator = MagicMock()
    mock_locator.first = mock_locator
    mock_locator.is_visible.return_value = False
    mock_page.locator.return_value = mock_locator

    # Mock Session for poll and download
    mock_sess = MagicMock()
    mock_sess.__enter__.return_value = mock_sess
    mock_sess.__exit__.return_value = None

    # Poll reports_list response
    mock_poll_resp = MagicMock()
    mock_poll_resp.text = (
        "<table><tr>"
        "<td>icon</td>"
        "<td>10/22/2026 1:16 AM WIB</td>"
        '<td><a href="/report_detail?id=9999">#UP time Report Switch Daily</a></td>'
        "<td>Custom</td><td>SW</td><td>bot</td>"
        "<td>10/19/2026 12:00 AM WIB</td>"
        "<td>10/19/2026 11:59 PM WIB</td>"
        "</tr></table>"
    )

    # Download zip response
    import io
    import zipfile

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        zf.writestr("top_cpu_utilization_9999.csv", "SW-01,12.0%\n")
        zf.writestr("avg_uptime_by_device_9999.csv", "99.0%,99.0%,99.0%,-\n")
    zip_bytes = zip_buffer.getvalue()

    mock_dl_resp = MagicMock()
    mock_dl_resp.iter_content.return_value = [zip_bytes]
    mock_dl_resp.raise_for_status.return_value = None

    def get_dispatch(url, *args, **kwargs):
        if "reports_list" in url:
            return mock_poll_resp
        elif "csv_export" in url:
            return mock_dl_resp
        return MagicMock()

    mock_sess.get.side_effect = get_dispatch

    out_dir = tmp_path / "raw" / "airwave" / "20261019"
    with (
        patch("network_report.collectors.airwave.sync_playwright") as mock_pw_start,
        patch("network_report.collectors.airwave.requests.Session", return_value=mock_sess),
    ):
        mock_pw_start.return_value.start.return_value = mock_playwright
        res_dir = client.fetch_report(date(2026, 10, 19), out_dir)

    assert res_dir == out_dir
    assert (out_dir / "top_cpu_utilization_9999.csv").exists()
    assert (out_dir / "avg_uptime_by_device_9999.csv").exists()


def test_airwave_stored_cookie_without_table_marker_triggers_fresh_login(tmp_path: Path):
    session_file = tmp_path / "storage_state.json"
    session_file.parent.mkdir(parents=True, exist_ok=True)
    session_file.write_text(
        json.dumps(
            {"cookies": [{"name": "MercuryAuthHandlerCookie_AMPAuth", "value": "expired-cookie"}]}
        )
    )

    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=session_file,
    )

    mock_sess = MagicMock()
    mock_sess.__enter__.return_value = mock_sess
    mock_sess.__exit__.return_value = None

    mock_check_resp = MagicMock()
    mock_check_resp.status_code = 200
    mock_check_resp.text = "<html><body>SPA shell without reports list table</body></html>"

    mock_post_resp = MagicMock()
    mock_post_resp.raise_for_status.return_value = None
    mock_cookie = MagicMock()
    mock_cookie.name = "MercuryAuthHandlerCookie_AMPAuth"
    mock_cookie.value = "fresh-cookie-12345"
    mock_sess.cookies = [mock_cookie]

    def mock_get(url, *args, **kwargs):
        if "reports_list" in url:
            return mock_check_resp
        return MagicMock()

    mock_sess.get.side_effect = mock_get
    mock_sess.post.return_value = mock_post_resp

    with patch("network_report.collectors.airwave.requests.Session", return_value=mock_sess):
        client.login()

    assert client._auth_cookie == "fresh-cookie-12345"
    mock_sess.post.assert_called_once()
    state = json.loads(session_file.read_text())
    assert state["cookies"][0]["value"] == "fresh-cookie-12345"


def test_airwave_stored_cookie_with_table_marker_is_reused(tmp_path: Path):
    session_file = tmp_path / "storage_state.json"
    session_file.parent.mkdir(parents=True, exist_ok=True)
    session_file.write_text(
        json.dumps(
            {
                "cookies": [
                    {"name": "MercuryAuthHandlerCookie_AMPAuth", "value": "valid-cookie-54321"}
                ]
            }
        )
    )

    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=session_file,
    )

    mock_sess = MagicMock()
    mock_sess.__enter__.return_value = mock_sess
    mock_sess.__exit__.return_value = None

    mock_check_resp = MagicMock()
    mock_check_resp.status_code = 200
    mock_check_resp.text = '<table class="reports_list"><tr><td><a href="/report_detail?id=99">Report</a></td></tr></table>'
    mock_sess.get.return_value = mock_check_resp

    with patch("network_report.collectors.airwave.requests.Session", return_value=mock_sess):
        client.login()

    assert client._auth_cookie == "valid-cookie-54321"
    mock_sess.get.assert_called_with("https://10.10.1.20/nf/reports_list", timeout=15)
    mock_sess.post.assert_not_called()


def test_airwave_login_force_always_performs_fresh_http_login(tmp_path: Path):
    session_file = tmp_path / "storage_state.json"
    session_file.parent.mkdir(parents=True, exist_ok=True)
    session_file.write_text(
        json.dumps(
            {"cookies": [{"name": "MercuryAuthHandlerCookie_AMPAuth", "value": "stored-cookie"}]}
        )
    )

    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=session_file,
    )

    mock_sess = MagicMock()
    mock_sess.__enter__.return_value = mock_sess
    mock_sess.__exit__.return_value = None

    mock_post_resp = MagicMock()
    mock_post_resp.raise_for_status.return_value = None
    mock_cookie = MagicMock()
    mock_cookie.name = "MercuryAuthHandlerCookie_AMPAuth"
    mock_cookie.value = "forced-fresh-cookie-999"
    mock_sess.cookies = [mock_cookie]
    mock_sess.post.return_value = mock_post_resp

    with patch("network_report.collectors.airwave.requests.Session", return_value=mock_sess):
        client.login(force=True)

    assert client._auth_cookie == "forced-fresh-cookie-999"
    mock_sess.post.assert_called_once()
    for call in mock_sess.get.call_args_list:
        args, _ = call
        if args:
            assert "reports_list" not in args[0]


def test_airwave_fetch_report_inline_dialog_triggers_forced_relogin(tmp_path: Path):
    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=tmp_path / "session.json",
        poll_interval_seconds=0,
        poll_timeout_minutes=1,
    )
    client._auth_cookie = "initial-cookie"

    mock_playwright = MagicMock()
    mock_browser = MagicMock()
    mock_context = MagicMock()
    mock_page = MagicMock()

    mock_playwright.firefox.launch.return_value = mock_browser
    mock_browser.new_context.return_value = mock_context
    mock_context.new_page.return_value = mock_page

    mock_login_user = MagicMock()
    mock_login_user.is_visible.side_effect = [True, False]

    mock_login_btn = MagicMock()
    mock_start_input = MagicMock()
    mock_end_input = MagicMock()
    mock_run_btn = MagicMock()
    default_locator = MagicMock()
    default_locator.first = default_locator

    def locator_dispatch(sel, *args, **kwargs):
        if sel == "#login-username-id":
            return mock_login_user
        if sel == "#login-button":
            return mock_login_btn
        if "start_string" in sel:
            return mock_start_input
        if "end_string" in sel:
            return mock_end_input
        if "definitions_do_and_run" in sel:
            return mock_run_btn
        return default_locator

    mock_page.locator.side_effect = locator_dispatch

    def fake_login(force=False):
        client._auth_cookie = "fresh-after-forced-relogin"

    client.login = MagicMock(side_effect=fake_login)
    client._collect_preexisting_ids = MagicMock(return_value=set())
    client._poll_report_id = MagicMock(return_value="7777")
    client._download_csv_zip = MagicMock(return_value=tmp_path / "dummy.zip")
    client._extract_zip = MagicMock(return_value=None)

    out_dir = tmp_path / "out"
    with patch("network_report.collectors.airwave.sync_playwright") as mock_pw_start:
        mock_pw_start.return_value.start.return_value = mock_playwright
        res = client.fetch_report(date(2026, 10, 19), out_dir)

    assert res == out_dir
    client.login.assert_called_once_with(force=True)
    mock_login_btn.click.assert_not_called()
    mock_context.add_cookies.assert_any_call(
        [
            {
                "name": "MercuryAuthHandlerCookie_AMPAuth",
                "value": "fresh-after-forced-relogin",
                "domain": "10.10.1.20",
                "path": "/",
            }
        ]
    )
    mock_context.set_extra_http_headers.assert_called_with(
        {"Cookie": "MercuryAuthHandlerCookie_AMPAuth=fresh-after-forced-relogin"}
    )
    mock_start_input.fill.assert_called_once()
    mock_end_input.fill.assert_called_once()
    mock_run_btn.click.assert_called_once()


def test_airwave_fetch_report_inline_dialog_fallback_failure_raises(tmp_path: Path):
    import pytest

    client = AirWaveClient(
        base_url="https://10.10.1.20",
        username="admin",
        password="secret",
        session_state_path=tmp_path / "session.json",
        max_retries=1,
    )
    client._auth_cookie = "initial-cookie"

    mock_playwright = MagicMock()
    mock_browser = MagicMock()
    mock_context = MagicMock()
    mock_page = MagicMock()

    mock_playwright.firefox.launch.return_value = mock_browser
    mock_browser.new_context.return_value = mock_context
    mock_context.new_page.return_value = mock_page

    mock_login_user = MagicMock()
    mock_login_user.is_visible.return_value = True

    mock_login_btn = MagicMock()
    default_locator = MagicMock()
    default_locator.first = default_locator

    def locator_dispatch(sel, *args, **kwargs):
        if sel == "#login-username-id":
            return mock_login_user
        if sel == "#login-button":
            return mock_login_btn
        return default_locator

    mock_page.locator.side_effect = locator_dispatch

    client.login = MagicMock(side_effect=RuntimeError("HTTP login failed"))
    client._collect_preexisting_ids = MagicMock(return_value=set())

    with (
        patch("network_report.collectors.airwave.sync_playwright") as mock_pw_start,
        pytest.raises(
            RuntimeError,
            match="AirWave authentication failed: inline login dialog remained visible",
        ),
    ):
        mock_pw_start.return_value.start.return_value = mock_playwright
        client.fetch_report(date(2026, 10, 19), tmp_path / "out")

    mock_login_btn.click.assert_called_once()
